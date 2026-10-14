"""Build synchronized Source Data from locked inputs and derived analysis views.

This builder runs before figure generation. Figure paths are planned outputs;
their existence and checksums are verified after rendering by the package audit.
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from analysis.submission_contract import (
    DERIVED, FIGURE_TITLES, MAIN_TITLES, ROOT, RUNS, STEMS,
    SUPP_STEMS, SUPP_TITLES, TITLE, prepare, read_csv, read_json, sha256,
)
from module1.bazil2016_ode import DEFAULT_PARAMS
from module1.neurmito_system import CARDIAC_IC, IDX
from scripts.source_audit.run_r010_robustness import initial_condition_variants

OUT = ROOT / "results" / "source_data"


def flatten_json(value, prefix=""):
    if isinstance(value, dict):
        for key, child in value.items():
            yield from flatten_json(child, f"{prefix}.{key}".strip("."))
    elif isinstance(value, list):
        for index, child in enumerate(value):
            yield from flatten_json(child, f"{prefix}[{index}]")
    else:
        yield {"json_path": prefix, "value": value, "value_type": type(value).__name__}


def scalar(value):
    if value is None:
        return None
    if isinstance(value, (int, float, bool)):
        return value
    text = str(value)
    if re.fullmatch(r"-?\d+", text):
        return int(text) if len(text.lstrip('-')) < 16 else text
    if re.fullmatch(r"-?(?:\d+\.?\d*|\.\d+)(?:[eE][+-]?\d+)?", text):
        return float(text)
    return text


def add_sheet(book, name, rows, columns=None):
    sheet = book.create_sheet(name)
    columns = columns or list(dict.fromkeys(key for row in rows for key in row))
    sheet.append(columns)
    for row in rows:
        sheet.append([scalar(row.get(key)) for key in columns])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2F5D7E")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for column in sheet.columns:
        letter = column[0].column_letter
        sheet.column_dimensions[letter].width = min(55, max(14, len(str(column[0].value)) + 2))
    return sheet


def parameter_rows():
    source = (ROOT / "module1/bazil2016_ode.py").read_text(encoding="utf-8")
    descriptions = {}
    for key, comment in re.findall(r"'([^']+)'\s*:\s*[^,]+,\s*#\s*([^\n]+)", source):
        descriptions[key] = comment.strip()
    rows = []
    for key, value in DEFAULT_PARAMS.items():
        if key.startswith("Vmax"):
            unit = "mM/s"
        elif key.startswith("Km") or key in {"Ki_H2O2", "O2_ext", "NAD_total", "Q_total", "Cc_total", "GSH_total", "FAD_total", "NADPH_target"}:
            unit = "mM"
        elif key.startswith("DPsi"):
            unit = "mV"
        elif key == "g_H":
            unit = "mM/(s mV)"
        elif key == "C_m":
            unit = "effective mM/mV, as required by implemented charge balance"
        elif key in {"k_NADPH_regen", "k_FAD_redox", "k_O2_in"}:
            unit = "1/s"
        else:
            unit = "dimensionless"
        rows.append({"parameter": key, "value": value, "unit": unit,
                     "description": descriptions.get(key, "Reduced-model implementation parameter"),
                     "provenance": "implemented reduced-model default; not a direct Bazil parameter copy"})
    for key, value, unit in [("k_Pi_in", .01, "1/s"), ("k_diff_H2O2", .01, "1/s"),
                              ("k_GPx_c", .005, "1/s"), ("redox_acceptor_half", .1, "mM"),
                              ("CS_OAA_half", .05, "mM"), ("SDH_SUC_half", .5, "mM")]:
        rows.append({"parameter": key, "value": value, "unit": unit,
                     "description": "Fixed constant in module1/neurmito_system.py",
                     "provenance": "implemented reduced model"})
    return rows


def initial_rows():
    variants = initial_condition_variants()
    return [{"state": name, "index": index, "unit": "mV" if name == "DPsi" else "mM",
             "default_ic": CARDIAC_IC[index],
             "low_energy_ic": variants["low_energy_reduced_ic"][index],
             "high_energy_ic": variants["high_energy_oxidized_ic"][index]}
            for name, index in sorted(IDX.items(), key=lambda item: item[1])]


RAW_SHEETS = {
    "R009/additive_proxy_residuals.csv": "Nonadditive_Residuals",
    "R009/mechanism_ablation_summary.csv": "Ablation_Summary",
    "R009/final_dpsi_by_variant_matrix.csv": "Ablation_Matrix",
    "R010/robustness_summary.csv": "Robustness_Summary",
    "R010/robustness_grid.csv": "Robustness_Grid",
    "R014B/resolved_transition_grid.csv": "Transition_Scan",
    "R014B/transition_band_summary.csv": "Transition_Bands",
    "R021/qc_summary.json": "Cohort_QC",
    "R021/conditions.csv": "Cohort_Conditions",
    "R021/events.csv": "Cohort_Events",
    "R021/split_manifest.csv": "Cohort_Splits",
    "R023/locked_test_metrics.json": "Locked_Metrics_JSON",
    "R023/negative_controls.json": "Negative_Controls",
    "R023/test_predictions.csv": "Test_Predictions",
    "R025/summary.csv": "Hybrid_Summary",
    "R025/metrics.csv": "Hybrid_Seed_Metrics",
    "R025/predictions.csv": "Hybrid_Predictions",
    "R025/metadata.json": "Hybrid_Metadata",
    "R042/digitization_uncertainty_propagation.json": "Digitization_Uncertainty",
    "R047/cross_source_formula_coordination.json": "Source_Definition_Audit",
    "R048/maximal_spare_restricted_mapping.json": "Restricted_Mapping",
    "R060/pmc4454361_control_relative_digitization.json": "R060_Extraction",
    "R060B/pmc5124357_control_relative_digitization.json": "R060B_Extraction",
    "R062/conditional_heterogeneity_mapping.json": "Conditional_Mapping",
}
DERIVED_SHEETS = ["Reserve_Observations", "EWS_Configurations", "EWS_Eligibility", "F4_Illustration",
                  "F5_Grid", "F5_Trajectories", "F6_Weights", "F6_Projections", "F6_Trajectories", "F6_Coverage"]


def panel_specs():
    main = {
        1: [["Source_Definition_Audit"], ["Source_Definition_Audit", "Reserve_Observations"],
            ["Reserve_Observations"], ["Reserve_Observations"], ["Reserve_Observations"],
            ["Restricted_Mapping"], ["Conditional_Mapping"], ["Conditional_Mapping"], ["Reserve_Observations"]],
        2: [["Transition_Scan"], ["Transition_Scan"], ["Transition_Scan"], ["Transition_Scan"],
            ["Transition_Bands"], ["Nonadditive_Residuals"], ["Ablation_Summary"],
            ["Robustness_Summary"], ["Robustness_Summary"]],
        3: [["Hybrid_Predictions"], ["Hybrid_Predictions"], ["Hybrid_Summary"], ["Hybrid_Summary"],
            ["Hybrid_Seed_Metrics"], ["Hybrid_Predictions"], ["Hybrid_Predictions"],
            ["Hybrid_Predictions"], ["Hybrid_Summary"]],
        4: [["F4_Illustration"], ["F4_Illustration"], ["F4_Illustration"], ["EWS_Configurations"],
            ["EWS_Configurations"], ["EWS_Configurations"], ["EWS_Configurations"],
            ["Negative_Controls"], ["Negative_Controls", "EWS_Configurations"], ["Test_Predictions"]],
        5: [["F5_Grid"], ["F5_Grid"], ["F5_Grid"], ["F5_Grid"], ["F5_Grid"],
            ["F5_Trajectories"], ["F5_Grid"], ["F5_Grid"], ["F5_Grid"], ["F5_Grid"]],
        6: [["F6_Weights"], ["F6_Trajectories"], ["F6_Projections"], ["F6_Weights"],
            ["F6_Projections"], ["F6_Weights"], ["F6_Weights"], ["F6_Coverage"], ["F6_Projections"]],
    }
    supp = {
        1: [["Reserve_Observations"], ["Source_Definition_Audit"], ["Digitization_Uncertainty"],
            ["Source_Definition_Audit", "Reserve_Observations"], ["Reserve_Observations"], []],
        2: [["Restricted_Mapping"], ["Conditional_Mapping"], ["Conditional_Mapping"],
            ["Restricted_Mapping"], ["Conditional_Mapping"], ["Conditional_Mapping"]],
        3: [["Hybrid_Summary"], ["Hybrid_Summary"], ["Hybrid_Seed_Metrics"],
            ["Hybrid_Predictions"], ["Hybrid_Predictions"], ["Hybrid_Predictions"]],
        4: [["Transition_Scan"], ["Transition_Scan"], ["Transition_Bands"],
            ["Nonadditive_Residuals"], ["Ablation_Summary"], ["Robustness_Summary"]],
        5: [["Cohort_Splits"], ["Cohort_QC"], ["Cohort_Events"], ["EWS_Configurations"],
            ["EWS_Configurations"], ["Negative_Controls"]],
        6: [[], ["Input_Inventory"], []],
    }
    return main, supp


def build_tables(inventory, summary):
    reserve = read_csv(DERIVED / "Reserve_Observations.csv")
    configs = read_csv(DERIVED / "EWS_Configurations.csv")
    tables = {
        "S1": [{"record": row["record"], "PMID": row["pmid"], "evidence_type": row["evidence_type"],
                "use": "source-specific respiratory evidence"} for row in reserve],
        "S2": [{key: row[key] for key in ["record", "pmid", "maximal_to_basal", "spare_to_basal", "spare_definition", "definition_difference"]} for row in reserve],
        "S3": [{key: row[key] for key in ["record", "source_file", "maximal_low", "maximal_high", "spare_low", "spare_high", "cross_source_equivalence"]} for row in reserve],
        "S4": parameter_rows(), "S5": initial_rows(),
        "S6": [{"figure": f"Figure {i}", "contribution": FIGURE_TITLES[i],
                "evidence_scope": "literature/source-local" if i == 1 else "reduced ODE" if i == 2 else
                                  "independent synthetic benchmark" if i == 3 else "retrospective event-aligned diagnostics" if i == 4 else "prescribed scenario",
                "restriction": "source spare definitions differ" if i == 1 else "model-defined boundary" if i == 2 else
                               "not biological training" if i == 3 else "different n/windows; future event anchors" if i == 4 else
                               "assumptions are not observed biological effects"}
               for i in range(1, 7)],
        "S7": read_csv(RUNS / "R025/metrics.csv"),
        "S8": read_csv(RUNS / "R014B/transition_band_summary.csv"),
        "S9": [{key: row[key] for key in ["model", "n", "positive_count", "negative_count", "excluded_count", "reference_time_s", "window_seconds", "lead_seconds", "threshold", "auroc", "auprc", "brier", "auroc_ci_low", "auroc_ci_high"]} for row in configs],
        "S10": inventory,
    }
    return tables


def main():
    summary = prepare()
    OUT.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    book.remove(book.active)
    add_sheet(book, "README", [
        {"field": "Title", "value": TITLE},
        {"field": "Panel counts", "value": "56 main + 33 supplementary = 89; one map row per lettered panel"},
        {"field": "Raw data", "value": "Locked inputs are copied unchanged; each numerical panel links to actual value tables"},
        {"field": "Evidence classes", "value": "literature-derived, reduced-model, synthetic benchmark, event-aligned diagnostic, prescribed scenario"},
        {"field": "Figure 4a-c", "value": "Deterministic illustration; not a sampled cohort trajectory"},
        {"field": "Figure 6c", "value": "Raw ROS weight, preset offset and plotted coordinate reported separately; biological/visual intent requires author confirmation"},
        {"field": "EWS exclusions", "value": "All effective n reproduced from 5001-point time grids and saved prediction group membership"},
        {"field": "EWS missing artifacts", "value": "R021 trajectories.csv and R022 lock/development artifacts absent; author confirmation required"},
    ])
    sources, inventory = {}, []
    for relpath, name in RAW_SHEETS.items():
        path = RUNS / relpath
        rows = read_csv(path) if path.suffix == ".csv" else list(flatten_json(read_json(path)))
        rows = [dict(row, source_file=str(path.relative_to(ROOT)).replace("\\", "/"), source_sha256=sha256(path)) for row in rows]
        add_sheet(book, name, rows)
        sources[name] = path
        inventory.append({"source_file": str(path.relative_to(ROOT)).replace("\\", "/"), "bytes": path.stat().st_size,
                          "sha256": sha256(path), "data_sheet": name, "status": "locked input unchanged"})
    for name in DERIVED_SHEETS:
        path = DERIVED / f"{name}.csv"
        add_sheet(book, name, read_csv(path))
        sources[name] = path
    add_sheet(book, "Model_Parameters", parameter_rows())
    add_sheet(book, "Model_Initial_States", initial_rows())
    add_sheet(book, "Input_Inventory", inventory)
    mappings = []
    specs = panel_specs()
    for supplementary, titles, panel_sources in [(False, MAIN_TITLES, specs[0]), (True, SUPP_TITLES, specs[1])]:
        for number, names in titles.items():
            for index, title in enumerate(names):
                data = panel_sources[number][index]
                numeric = bool(data)
                paths = [sources[name] for name in data if name in sources]
                if "Input_Inventory" in data:
                    paths = [ROOT / "analysis/submission_contract.py"]
                if not paths:
                    paths = [ROOT / "scripts/r071/build_r071_supplementary.py"]
                function = ({1: "source_audit_s1", 2: "mapping_s2", 3: "hybrid_s3", 4: "phase_s4", 5: "ews_s5", 6: "reproducibility_s6"}[number]
                            if supplementary else {1: "figure1", 2: "figure3", 3: "figure2", 4: "figure4", 5: "figure5", 6: "figure6"}[number])
                generator = f"scripts/r071/build_r071_{'supplementary' if supplementary else 'all'}.py::{function}"
                stem = SUPP_STEMS[number] if supplementary else STEMS[number]
                mappings.append({"figure_id": f"Supplementary Figure S{number}" if supplementary else f"Figure {number}",
                                 "panel_id": chr(97 + index), "panel_title": title,
                                 "panel_type": "schematic" if not numeric else "illustrative" if not supplementary and number == 4 and index < 3 else "data",
                                 "source_data_status": "numeric" if numeric else "schematic/no numeric source data",
                                 "data_sheets": ";".join(data),
                                 "data_ranges": ";".join(f"{name}!{book[name].dimensions}" for name in data),
                                 "source_files": ";".join(str(path.relative_to(ROOT)).replace("\\", "/") for path in paths),
                                 "source_sha256": ";".join(sha256(path) for path in paths),
                                 "generator": generator,
                                 "figure_file": f"results/{'supplementary_figures' if supplementary else 'figures'}/{stem}.pdf",
                                 "evidence_scope": "see S6 and README; retain source/formula/analysis-unit restrictions"})
    add_sheet(book, "Panel_Map", mappings)
    add_sheet(book, "Panel_QA", [{"figure": row["figure_id"], "panel": row["panel_id"],
                                 "value_tables_present": bool(row["data_sheets"]) or row["panel_type"] == "schematic",
                                 "source_files_present": True, "figure_validation": "performed after regeneration"}
                                for row in mappings])
    dictionary = [
        {"variable": "DPsi", "unit": "mV", "definition": "implemented membrane potential"},
        {"variable": "ATP_m, H2O2_m and concentration states", "unit": "mM", "definition": "reduced-model state concentrations"},
        {"variable": "rotenone_inhibition", "unit": "fraction", "definition": "model operator; figure axes display 100 times fraction"},
        {"variable": "paraquat_index", "unit": "dimensionless", "definition": "model ROS amplification index"},
        {"variable": "source spare/basal", "unit": "dimensionless", "definition": "algebraic for R033/R039; independently digitized reported spare for R060/R060B"},
        {"variable": "lead_seconds", "unit": "s", "definition": "prescribed distance before observed event; retrospective"},
        {"variable": "rolling_variance", "unit": "mV squared", "definition": "sample variance, ddof=1"},
        {"variable": "dpsi_slope", "unit": "mV per sample index", "definition": "linear fit against sample index, not seconds"},
        {"variable": "F6 plotted_ros_coordinate", "unit": "uncalibrated coordinate", "definition": "specified ROS weight + 0.08 times context index"},
        {"variable": "scenario score range", "unit": "dimensionless", "definition": "prescribed +/-0.08; not statistical error"},
    ]
    add_sheet(book, "Variable_Dictionary", dictionary)
    book.save(OUT / "NeurMito_Source_Data.xlsx")
    tables = build_tables(inventory, summary)
    table_book = Workbook()
    table_book.remove(table_book.active)
    for name, rows in tables.items():
        add_sheet(table_book, name, rows)
    table_book.save(OUT / "Supplementary_Tables_S1-S10.xlsx")
    (OUT / "supplementary_tables.json").write_text(json.dumps(tables, indent=2), encoding="utf-8")
    contract = {"title": TITLE, "panels": mappings, "figure_titles": FIGURE_TITLES,
                "figure_stems": STEMS, "supplementary_stems": SUPP_STEMS,
                "panel_count": len(mappings), "main_panels": 56, "supplementary_panels": 33}
    (OUT / "submission_contract.json").write_text(json.dumps(contract, indent=2), encoding="utf-8")
    print(json.dumps({"source_data": str(OUT / "NeurMito_Source_Data.xlsx"), "panels": len(mappings), "sheets": len(book.sheetnames)}, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
